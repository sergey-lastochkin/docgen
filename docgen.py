#!/usr/bin/env python3
"""docgen.py — комплект документов по заказу: смета, договор, ТЗ (приложение № 1), акт.

  python3 docgen.py specs/<заказ>.json            -> out/<номер>/*.docx
  python3 docgen.py specs/<заказ>.json --check    -> только проверка спецификации

Реквизиты исполнителя — executor.json (один раз). Всё про заказ — в спецификации.
Формат: Times New Roman 12, интервал 1,15, поля 2/1 см.
"""
from __future__ import annotations

import datetime as dt
import json
import pathlib
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = pathlib.Path(__file__).resolve().parent
SIGNATURE = HERE / 'signature.png'
USE_SIG = True  # выставляется из spec['facsimile']
MONTHS = ['января', 'февраля', 'марта', 'апреля', 'мая', 'июня', 'июля', 'августа',
          'сентября', 'октября', 'ноября', 'декабря']

# ---------- сумма прописью ----------
_UNITS = {
    'm': ['', 'один', 'два', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять'],
    'f': ['', 'одна', 'две', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять'],
}
_TEENS = ['десять', 'одиннадцать', 'двенадцать', 'тринадцать', 'четырнадцать', 'пятнадцать',
          'шестнадцать', 'семнадцать', 'восемнадцать', 'девятнадцать']
_TENS = ['', '', 'двадцать', 'тридцать', 'сорок', 'пятьдесят', 'шестьдесят', 'семьдесят',
         'восемьдесят', 'девяносто']
_HUNDREDS = ['', 'сто', 'двести', 'триста', 'четыреста', 'пятьсот', 'шестьсот', 'семьсот',
             'восемьсот', 'девятьсот']


def _plural(n: int, forms: tuple[str, str, str]) -> str:
    n = abs(n) % 100
    if 11 <= n <= 19:
        return forms[2]
    n %= 10
    return forms[0] if n == 1 else forms[1] if 2 <= n <= 4 else forms[2]


def _triad(n: int, gender: str) -> list[str]:
    words = [_HUNDREDS[n // 100]]
    t = n % 100
    if 10 <= t <= 19:
        words.append(_TEENS[t - 10])
    else:
        words += [_TENS[t // 10], _UNITS[gender][t % 10]]
    return [w for w in words if w]


def rubles_in_words(amount: int) -> str:
    """28000 -> 'Двадцать восемь тысяч' (только целые рубли, до миллиардов)."""
    if amount == 0:
        return 'Ноль'
    parts = []
    millions, rest = divmod(amount, 1_000_000)
    thousands, units = divmod(rest, 1000)
    if millions:
        parts += _triad(millions, 'm') + [_plural(millions, ('миллион', 'миллиона', 'миллионов'))]
    if thousands:
        parts += _triad(thousands, 'f') + [_plural(thousands, ('тысяча', 'тысячи', 'тысяч'))]
    parts += _triad(units, 'm')
    s = ' '.join(parts)
    return s[0].upper() + s[1:]


def money(amount: int) -> str:
    return f'{amount:,}'.replace(',', ' ')


def money_full(amount: int) -> str:
    return f'{money(amount)} ({rubles_in_words(amount)}) рублей 00 копеек'


def date_long(iso: str) -> str:
    d = dt.date.fromisoformat(iso)
    return f'{d.day:02d} {MONTHS[d.month - 1]} {d.year} г.'


def date_short(iso: str) -> str:
    return dt.date.fromisoformat(iso).strftime('%d.%m.%Y')


# ---------- оформление ----------
class Doc:
    def __init__(self):
        self.d = Document()
        for s in self.d.sections:
            s.left_margin, s.right_margin = Cm(2), Cm(1)
            s.top_margin, s.bottom_margin = Cm(2), Cm(2)
        st = self.d.styles['Normal']
        st.font.name = 'Times New Roman'
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        st.font.size = Pt(12)

    def p(self, text='', size=12, bold=False, align='justify', space_after=4, indent=False):
        par = self.d.add_paragraph()
        par.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
        pf = par.paragraph_format
        pf.line_spacing = 1.15
        pf.space_after = Pt(space_after)
        pf.space_before = Pt(0)
        if indent:
            pf.first_line_indent = Cm(1.0)
        if text:
            r = par.add_run(text)
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor(0, 0, 0)
        return par

    def title(self, text, size=14):
        self.p(text, size=size, bold=True, align='center', space_after=2)

    def section(self, text):
        par = self.p(text, bold=True, align='center', space_after=4)
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.keep_with_next = True

    def table(self, rows, widths, header=True, size=11, borders=True):
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        if borders:
            t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c)
                cell.width = Cm(widths[c])
                cell.text = ''
                lines = str(val).split('\n')
                for i, line in enumerate(lines):
                    par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                    par.paragraph_format.line_spacing = 1.0
                    par.paragraph_format.space_after = Pt(0)
                    if line == '{SIG}':
                        if SIGNATURE.exists():
                            par.add_run().add_picture(str(SIGNATURE), width=Cm(3.0))
                        continue
                    run = par.add_run(line)
                    run.font.size = Pt(size)
                    run.font.bold = header and r == 0
        self.p(space_after=2)
        return t

    def save(self, path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(path)
        return path


# ---------- стороны ----------
def client_intro(c):
    if c['type'] == 'ИП':
        return (f"Индивидуальный предприниматель {c['full_name']}, ОГРНИП {c['ogrn']}")
    return (f"{c['org_form_full']} «{c['name']}», ОГРН {c['ogrn']}, в лице {c['signer_position_gen']} "
            f"{c['signer_full_gen']}, {c.get('signer_acting', 'действующего')} на основании {c['signer_basis']}")


def named(c):
    return 'именуемый' if c['type'] == 'ИП' else c.get('named', 'именуемое')


def client_short(c):
    return f"ИП {c['full_name']}" if c['type'] == 'ИП' else f"{c['org_form']} «{c['name']}»"


def executor_intro(e):
    return (f"{e['full_name']}, ИНН {e['inn']}, физическое лицо, применяющее специальный налоговый режим "
            '«Налог на профессиональный доход» (НПД)')


def requisites_block(c):
    lines = [client_short(c)]
    if c.get('inn'):
        lines.append(f"ИНН {c['inn']}")
    if c.get('kpp'):
        lines.append(f"КПП {c['kpp']}")
    lines.append(f"{'ОГРНИП' if c['type'] == 'ИП' else 'ОГРН'} {c['ogrn']}")
    if c.get('address'):
        lines.append(c['address'])
    if c.get('account'):
        lines += [f"р/с {c['account']}", c['bank'], f"БИК {c['bik']}", f"к/с {c['ks']}"]
    return '\n'.join(lines)


def executor_requisites(e):
    lines = [e['full_name'], 'Плательщик НПД', f"ИНН {e['inn']}", f"р/с {e['account']}", e['bank'],
             f"БИК {e['bik']}", f"к/с {e['ks']}"]
    if e.get('bank_inn'):
        lines.append(f"ИНН/КПП банка {e['bank_inn']}/{e.get('bank_kpp', '')}")
    lines += [x for x in (e.get('phone'), e.get('email'), e.get('telegram')) if x]
    return '\n'.join(lines)


def signatures(doc, c, e, with_dates=False):
    left = f"ЗАКАЗЧИК\n\n________________ / {c['signer_short']} /"
    right = f"ИСПОЛНИТЕЛЬ\n{'{SIG}' if USE_SIG else ''}\n________________ / {e['signer_short']} /"
    if with_dates:
        left += '\nДата: «___» __________ 20__ г.'
        right += '\nДата: «___» __________ 20__ г.'
    doc.table([[left, right]], [8.5, 8.5], header=False, size=12, borders=False)


# ---------- документы ----------
def estimate(spec, e):
    c = spec['client']
    doc = Doc()
    doc.title('СМЕТА')
    doc.p(spec['estimate_title'], size=13, align='center', space_after=8)
    doc.table([
        ['Дата составления:', date_short(spec['date'])],
        ['Заказчик:', client_short(c) if spec.get('estimate_show_client', True) else 'ИП ______________________________'],
        ['Исполнитель:', e['full_name']],
        ['ИНН:', e['inn']],
        ['Системы:', spec['systems']],
    ], [4.5, 12.5], header=False, size=11, borders=False)
    doc.p('Исходные данные: ' + spec['initial_data'], size=11)
    rows = [['№', 'Наименование работ', 'Состав / результат', 'Стоимость, руб.']]
    for i, it in enumerate(spec['items'], 1):
        rows.append([str(i), it['name'], it['result'], money(it['price'])])
    doc.table(rows, [1.0, 4.5, 9.0, 2.5], size=10)
    doc.p(f"ИТОГО: {money_full(total(spec))}", bold=True, align='right')
    doc.p('Условия выполнения работ', bold=True)
    pp = spec.get('prepay_pct', 50)
    terms = [f"Срок выполнения: {spec['estimate_term']}.",
             f'Оплата: {pp} % предоплата, {100 - pp} % после демонстрации и приемки результата.']
    terms += spec.get('estimate_notes', [])
    terms.append(f"Смета действительна {spec.get('estimate_valid_days', 7)} календарных дней с даты составления.")
    for t in terms:
        doc.p('— ' + t, size=10.5)
    if spec.get('estimate_reference'):
        doc.p('Справочно: ' + spec['estimate_reference'], size=10.5)
    signatures(doc, c, e, with_dates=True)
    return doc


def contract(spec, e):
    c = spec['client']
    num, tot = spec['number'], total(spec)
    pp = spec.get('prepay_pct', 50)
    adv = round(tot * pp / 100)
    doc = Doc()
    doc.title('ДОГОВОР ВОЗМЕЗДНОГО ОКАЗАНИЯ УСЛУГ')
    doc.p(f'№ {num}', align='center')
    doc.table([[spec.get('city', 'г. Санкт-Петербург'), date_long(spec['date'])]], [8.5, 8.5],
              header=False, size=12, borders=False)
    doc.p(f"{client_intro(c)}, {named(c)} в дальнейшем «Заказчик», с одной стороны, и {executor_intro(e)}, "
          'именуемый в дальнейшем «Исполнитель», с другой стороны, совместно именуемые «Стороны», '
          'заключили настоящий договор о нижеследующем.', indent=True)

    def clauses(head, items):
        doc.section(head)
        for it in items:
            doc.p(it)

    clauses('1. ПРЕДМЕТ ДОГОВОРА', [
        f"1.1. Исполнитель обязуется по заданию Заказчика оказать услуги по {spec['subject']}, "
        'а Заказчик обязуется принять и оплатить услуги.',
        '1.2. Состав, объем, этапы работ, критерии приемки и ограничения определены в Приложении № 1 '
        '(Техническое задание), являющемся неотъемлемой частью настоящего договора.',
        '1.3. ' + spec.get('subject_scope_note',
                           'Услуги оказываются в отношении уже используемой Заказчиком инфраструктуры и '
                           'существующих учетных записей. Разработка отдельного программного продукта или '
                           'передача исключительных прав по настоящему договору не предполагается, если иное '
                           'письменно не согласовано Сторонами.'),
    ])
    start, end = spec['start'], spec['end']
    days = (dt.date.fromisoformat(end) - dt.date.fromisoformat(start)).days + 1
    clauses('2. ПОРЯДОК И СРОК ОКАЗАНИЯ УСЛУГ', [
        f"2.1. Исполнитель приступает к оказанию услуг {date_long(start)[:-3]} при условии получения "
        'авансового платежа и предоставления Заказчиком необходимых доступов, исходных данных и примеров '
        'документов/операций не позднее указанной даты.',
        f"2.2. Конечный срок завершения и передачи результата Заказчику - не позднее {date_long(end)[:-3]} "
        f"включительно. Период с {date_short(start)} по {date_short(end)} является согласованным сроком "
        f'выполнения работ ({days} календарных дней).',
        '2.3. Изменение срока, установленного пунктом 2.2, допускается только по письменному согласованию '
        'Сторон при наличии объективных обстоятельств, препятствующих выполнению работ: задержки со стороны '
        'Заказчика в предоставлении доступов, исходных данных или необходимых действий; недоступности либо '
        'существенного изменения работы сторонних сервисов; выявления технического ограничения, которое '
        'нельзя было разумно установить до начала работ. Исполнитель обязан без неоправданной задержки '
        'сообщить Заказчику о таком обстоятельстве, описать его влияние на срок и продолжать выполнение '
        'доступной части работ.',
        '2.4. Исполнитель вправе выполнять настройку на тестовых данных и/или на ограниченном наборе реальных '
        'операций. Действия, требующие использования квалифицированной электронной подписи или подписания '
        'юридически значимых документов, выполняются Заказчиком самостоятельно по инструкции Исполнителя. '
        'Закрытые ключи КЭП Исполнителю не передаются.',
    ])
    clauses('3. СТОИМОСТЬ УСЛУГ И РАСЧЕТЫ', [
        f'3.1. Общая стоимость услуг по настоящему договору составляет {money_full(tot)}. НДС не '
        'предъявляется в связи с применением Исполнителем специального налогового режима НПД.',
        '3.2. ' + spec['payment_text'] if spec.get('payment_text') else
        f'3.2. Оплата производится в два этапа: аванс {pp}% в размере {money_full(adv)} - в течение 2 (двух) '
        f'рабочих дней после подписания договора; оставшиеся {100 - pp}% в размере {money_full(tot - adv)} - '
        'в течение 3 (трех) рабочих дней после подписания Акта оказанных услуг либо наступления приемки в '
        'порядке пункта 4.3.',
        '3.3. Стоимость подписок, тарифов, комиссий и иных платежей третьим лицам в стоимость услуг не '
        'входит и оплачивается Заказчиком самостоятельно.',
        '3.4. Работы, не предусмотренные Приложением № 1, выполняются только после отдельного согласования '
        'объема, сроков и стоимости.',
    ])
    clauses('4. СДАЧА И ПРИЕМКА УСЛУГ', [
        '4.1. После завершения услуг Исполнитель направляет Заказчику Акт оказанных услуг и сообщает о '
        'готовности результата к проверке.',
        f'4.2. Заказчик обязуется проверить переданный результат и не позднее {date_short(end)} подписать Акт '
        'либо направить Исполнителю мотивированный перечень конкретных несоответствий Техническому заданию.',
        '4.3. Если результат передан в срок и Заказчик в течение 3 (трех) рабочих дней после его передачи не '
        'направил подписанный Акт либо мотивированные письменные возражения, услуги считаются принятыми.',
        '4.4. Замечания, относящиеся к согласованному объему услуг, Исполнитель устраняет без дополнительной '
        'оплаты в разумный срок. Новые требования и изменение бизнес-процесса после согласования Технического '
        'задания являются дополнительными работами.',
    ])
    duties = [
        '5.1. Исполнитель обязуется оказать услуги добросовестно, использовать предоставленные доступы только '
        'для целей настоящего договора, не раскрывать полученные учетные данные третьим лицам и по завершении '
        'работ прекратить их использование.',
        '5.2. Заказчик обязуется своевременно предоставить корректные доступы к системам и учетным записям, '
        'тестовые данные, примеры документов и иную информацию, необходимую для выполнения Технического '
        'задания.',
    ]
    n = 3
    for extra in spec.get('extra_duties', []):
        duties.append(f'5.{n}. {extra}')
        n += 1
    duties.append(f'5.{n}. Заказчик самостоятельно принимает решения по хозяйственным операциям, подписывает '
                  'юридически значимые документы и несет ответственность за достоверность предоставленных данных.')
    clauses('5. ОБЯЗАННОСТИ СТОРОН', duties)
    clauses('6. НПД И ЧЕК', [
        '6.1. Исполнитель подтверждает, что на дату заключения договора состоит на учете как плательщик НПД '
        'и самостоятельно исполняет обязанности по уплате налога на профессиональный доход.',
        '6.2. По каждому полученному платежу Исполнитель формирует чек в установленном порядке и передает его '
        'Заказчику. В случае утраты права на применение НПД Исполнитель обязуется уведомить Заказчика без '
        'неоправданной задержки.',
    ])
    pen, cap = spec.get('penalty_pct'), spec.get('penalty_cap_pct', 20)
    resp = [
        '7.1. Стороны несут ответственность за неисполнение обязательств в соответствии с законодательством '
        'Российской Федерации и настоящим договором.',
        '7.2. Исполнитель не отвечает за сбои, блокировки, изменения правил, тарифов, API и функциональности '
        'сторонних сервисов, а также за отказ таких сервисов принять данные при корректно выполненной со '
        'стороны Исполнителя настройке.',
        '7.3. ' + spec.get('clause_7_3', 'Исполнитель не оказывает в рамках настоящего договора юридические, '
                           'налоговые или бухгалтерские консультации.'),
        '7.4. Если иное прямо не установлено законом, совокупная ответственность Исполнителя ограничивается '
        'суммой фактически уплаченного Заказчиком вознаграждения по договору.',
    ]
    if pen:
        resp.append(f'7.5. При нарушении конечного срока, установленного пунктом 2.2, по причинам, зависящим '
                    f'от Исполнителя, Заказчик вправе потребовать уплаты пени в размере {pen}% от общей '
                    f'стоимости услуг за каждый календарный день просрочки, но не более {cap}% от общей '
                    'стоимости услуг. Пеня не начисляется за период переноса срока в порядке пункта 2.3, а '
                    'также за задержки, за которые Исполнитель не несет ответственности.')
    clauses('7. ОТВЕТСТВЕННОСТЬ И ОГРАНИЧЕНИЯ', resp)
    clauses('8. КОНФИДЕНЦИАЛЬНОСТЬ', [
        '8.1. Стороны обязуются не раскрывать третьим лицам полученные в связи с исполнением договора пароли, '
        'токены, коммерческую информацию, персональные данные и иные сведения ограниченного доступа, кроме '
        'случаев, предусмотренных законом или необходимых для исполнения договора.',
        '8.2. Исполнитель не размещает реальные реквизиты Заказчика, ключи доступа и персональные данные в '
        'публичных репозиториях и общедоступных сервисах.',
    ])
    clauses('9. ПРОЧИЕ УСЛОВИЯ', [
        '9.1. Все изменения объема услуг и дополнительные работы согласуются Сторонами в письменной форме, в '
        'том числе путем обмена сообщениями по согласованным каналам связи.',
        '9.2. Скан-копии подписанных документов и документы, подписанные электронной подписью, признаются '
        'Сторонами допустимым способом оформления взаимоотношений.'
        + (' Стороны допускают использование факсимильного воспроизведения подписи Исполнителя на '
           'договоре, приложениях к нему, счетах и актах (п. 2 ст. 160 ГК РФ).' if USE_SIG else ''),
        '9.3. Споры Стороны стремятся урегулировать путем переговоров и письменной претензии, при '
        'недостижении соглашения — в порядке, установленном законодательством Российской Федерации.',
        '9.4. Договор вступает в силу с даты его подписания обеими Сторонами и действует до полного '
        'исполнения обязательств.',
        '9.5. Приложение № 1 и Акт оказанных услуг являются неотъемлемыми документами к настоящему договору.',
    ])
    doc.section('10. РЕКВИЗИТЫ И ПОДПИСИ СТОРОН')
    doc.table([[f"ЗАКАЗЧИК\n{requisites_block(c)}\n\n________________ / {c['signer_short']} /",
                f"ИСПОЛНИТЕЛЬ\n{executor_requisites(e)}\n{'{SIG}' if USE_SIG else ''}\n________________ / {e['signer_short']} /"]],
              [8.5, 8.5], header=False, size=11, borders=False)
    return doc


def tz(spec, e):
    c = spec['client']
    doc = Doc()
    doc.p(f"Приложение № 1 к Договору № {spec['number']} от {date_long(spec['date'])[:-3]}", align='right')
    doc.title('ТЕХНИЧЕСКОЕ ЗАДАНИЕ')
    doc.p(spec['tz_title'], size=13, align='center', space_after=8)

    def numbered(head, num, items):
        doc.section(head)
        for i, it in enumerate(items, 1):
            doc.p(f'{num}.{i}. {it}')

    numbered('1. ИСХОДНАЯ СИТУАЦИЯ', 1, spec['situation'])
    numbered('2. СОСТАВ УСЛУГ', 2, spec['services'])
    numbered('3. КРИТЕРИИ ПРИЕМКИ', 3, spec['acceptance'])
    numbered('4. НЕ ВХОДИТ В СТОИМОСТЬ', 4, spec['excluded'])
    numbered('5. СТОИМОСТЬ И СРОК', 5, [
        f'Стоимость услуг по настоящему Техническому заданию включена в цену Договора и составляет '
        f'{money_full(total(spec))}.',
        f"Работы выполняются в период с {date_long(spec['start'])[:-3]} по {date_long(spec['end'])[:-3]} "
        f"включительно. Конечный срок передачи результата - {date_long(spec['end'])[:-3]}. Перенос срока "
        'допускается только по письменному согласованию Сторон в случаях и порядке, предусмотренных '
        'пунктом 2.3 Договора.',
    ])
    doc.section('6. ПОДПИСИ СТОРОН')
    signatures(doc, c, e)
    return doc


def act(spec, e):
    c = spec['client']
    tot = total(spec)
    doc = Doc()
    doc.title('АКТ ОКАЗАННЫХ УСЛУГ')
    doc.p(f"к Договору № {spec['number']} от {date_long(spec['date'])[:-3]}", align='center')
    doc.table([[spec.get('city', 'г. Санкт-Петербург'), '«___» __________ 20__ г.']], [8.5, 8.5],
              header=False, size=12, borders=False)
    doc.p(f"{client_intro(c)}, {named(c)} «Заказчик», и {e['full_name']}, ИНН {e['inn']}, плательщик НПД, "
          'именуемый «Исполнитель», составили настоящий Акт о нижеследующем:', indent=True)
    for i, t in enumerate([
        f"Исполнитель оказал, а Заказчик принял услуги по {spec['subject']} в соответствии с Договором и "
        'Приложением № 1.',
        f'Стоимость оказанных услуг составляет {money_full(tot)}. НДС не предъявляется в связи с применением '
        'Исполнителем НПД.',
        'Услуги оказаны в согласованном объеме. Заказчик к качеству, объему и срокам оказанных услуг претензий '
        'не имеет.',
        'Настоящий Акт является основанием для окончательного расчета по Договору. После получения '
        'соответствующего платежа Исполнитель формирует и передает Заказчику чек НПД.',
    ], 1):
        doc.p(f'{i}. {t}')
    signatures(doc, c, e)
    return doc


def total(spec):
    return sum(int(it['price']) for it in spec['items'])


REQUIRED = ['number', 'date', 'start', 'end', 'client', 'subject', 'systems', 'estimate_title',
            'estimate_term', 'initial_data', 'items', 'tz_title', 'situation', 'services', 'acceptance',
            'excluded']


def check(spec):
    errors = [f'нет поля {k}' for k in REQUIRED if k not in spec]
    c = spec.get('client', {})
    for k in ['type', 'ogrn', 'signer_short'] + (['full_name'] if c.get('type') == 'ИП' else
                                                  ['name', 'org_form', 'org_form_full', 'signer_position_gen',
                                                   'signer_full_gen', 'signer_basis']):
        if not c.get(k):
            errors.append(f'клиент: нет {k}')
    for k in ('date', 'start', 'end'):
        try:
            dt.date.fromisoformat(spec.get(k, ''))
        except ValueError:
            errors.append(f'{k}: дата в формате ГГГГ-ММ-ДД')
    if not errors and spec['start'] > spec['end']:
        errors.append('start позже end')
    for i, it in enumerate(spec.get('items', []), 1):
        if not all(k in it for k in ('name', 'result', 'price')):
            errors.append(f'работа {i}: нужны name, result, price')
    return errors


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 2
    spec_path = pathlib.Path(argv[1])
    spec = json.loads(spec_path.read_text(encoding='utf-8'))
    errors = check(spec)
    if errors:
        print('Ошибки в спецификации:\n- ' + '\n- '.join(errors))
        return 1
    print(f"итого {money_full(total(spec))}")
    if '--check' in argv:
        return 0
    e = json.loads((HERE / 'executor.json').read_text(encoding='utf-8'))
    global USE_SIG
    USE_SIG = spec.get('facsimile', True) and SIGNATURE.exists()
    out = HERE / 'out' / spec['number'].replace('/', '-')
    surname = spec['client'].get('file_suffix', '')
    for name, builder in [('Смета', estimate), ('Договор', contract), ('ТЗ', tz), ('Акт', act)]:
        path = builder(spec, e).save(out / f'{name}{(" " + surname) if surname else ""}.docx')
        print('готово:', path)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
